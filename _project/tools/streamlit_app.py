import streamlit as st
import json
import pandas as pd
import os
import re
import sqlite3
from collections import defaultdict
from fpdf import FPDF
from docx import Document
from gpt_handler import analyze_long_text as ana
from rule_engine import apply_rules_to_chart, load_rules, analyze_with_rules
from visualizer import show_visualizations

# === 설정 ===
st.set_page_config(page_title="통합 분석기", layout="wide")
st.title("📊 통합 분석기")
st.caption("AI + 규칙 기반 분석 + 키워드 추출 + PDF 리포트")

# === 유틸 함수 ===
def extract_text(file):
    suffix = file.name.split(".")[-1]
    if suffix == "docx":
        doc = Document(file)
        return "\n".join(p.text for p in doc.paragraphs)
    else:
        return file.read().decode("utf-8")

def extract_rules_advanced(text):
    pattern = re.compile(r'(.+?)\s*(경우|일 경우|이라면|이면|일 때|인 경우|했을 때|하게 되면|한다면|할 경우|시|때)\s*[:,]?\s*(.+?)(?:\.|$)')
    rules = {}
    for match in pattern.finditer(text):
        condition = match.group(1).strip()
        result = match.group(3).strip()
        condition_clean = re.sub(r"(은는이가|을를|에|에서|의|로|으로|도|만|과|와|이며|이나|하고|부터|까지)$", "", condition)
        result_clean = re.sub(r"(은는이가|을를|에|에서|의|로|으로|도|만|과|와|이며|이나|하고|부터|까지)$", "", result)
        rules.setdefault(condition_clean, []).append(result_clean)
    return rules

def extract_sections(text):
    keywords = ['혼인', '결혼', '再婚', '初婚', '배우자궁', '부궁', '夫星', '妻宮', '副夫宮']
    lines = text.splitlines()
    results, current_block, capture = [], [], False
    for line in lines:
        if any(kw in line for kw in keywords):
            capture = True
        if capture:
            if line.strip():
                current_block.append(line.strip())
            else:
                if current_block:
                    results.append("\n".join(current_block))
                    current_block, capture = [], False
    return results, keywords

def analyze_keywords(results, keywords):
    keyword_hits = defaultdict(list)
    for text in results:
        for kw in keywords:
            if kw in text:
                keyword_hits[kw].append(text)
    df = {"키워드": [], "빈도": [], "예시 일부": []}
    for k, v in keyword_hits.items():
        df["키워드"].append(k)
        df["빈도"].append(len(v))
        df["예시 일부"].append(v[0][:100] + "..." if v else "")
    return pd.DataFrame(df), keyword_hits

def generate_pdf_report(filename, rule_data: dict, keyword_summary_df):
    pdf = FPDF()
    pdf.add_page()
    font_path = "NanumGothic.ttf"
    pdf.add_font("Nanum", "", font_path, uni=True)
    pdf.add_font("Nanum", "B", font_path, uni=True)
    pdf.set_font("Nanum", 'B', 16)
    pdf.cell(200, 10, txt="사주 문서 분석 보고서", ln=True, align="C")
    pdf.ln(10)
    pdf.set_font("Nanum", 'B', 12)
    pdf.cell(200, 10, txt="추출된 해석 규칙", ln=True)
    for k, v_list in rule_data.items():
        pdf.set_font("Nanum", 'B', 11)
        pdf.cell(200, 8, txt=f"[{k}]", ln=True)
        pdf.set_font("Nanum", '', 11)
        for v in v_list:
            pdf.multi_cell(0, 6, f"- {v}")
    pdf.add_page()
    pdf.set_font("Nanum", 'B', 12)
    pdf.cell(200, 10, txt="혼인 키워드 분석 요약", ln=True)
    for i, row in keyword_summary_df.iterrows():
        pdf.set_font("Nanum", '', 11)
        line = f"{row['키워드']} ({row['빈도']}회): {row['예시 일부']}"
        pdf.multi_cell(0, 6, line)
    pdf.output(filename)

# === 파일 업로드 및 처리 ===
uploaded_file = st.file_uploader("📂 파일 업로드 (json, txt, docx)", type=["json", "txt", "docx"])
if uploaded_file:
    file_ext = uploaded_file.name.split(".")[-1]
    if file_ext == "json":
        saju_data = json.load(uploaded_file)
        text = json.dumps(saju_data, ensure_ascii=False, indent=2)
    else:
        text = extract_text(uploaded_file)

    st.success("✅ 파일 처리 완료")
    st.text_area("📑 텍스트 미리보기", text, height=300)

    col1, col2, col3 = st.columns(3)

    with col1:
        st.markdown("### 🤖 GPT 분석")
        if st.button("GPT로 분석"):
            with st.spinner("GPT 분석 중..."):
                result = ana(text)
                st.text_area("📘 분석 결과", result, height=300)

    with col2:
        st.markdown("### 📜 규칙 기반 추출")
        if st.button("규칙 분석 실행"):
            with st.spinner("규칙 분석 중..."):
                rules = extract_rules_advanced(text)
                st.json(rules)

    with col3:
        st.markdown("### 🔍 키워드 분석")
        if st.button("혼인 키워드 분석"):
            with st.spinner("혼인 키워드 분석 중..."):
                blocks, keywords = extract_sections(text)
                df, keyword_map = analyze_keywords(blocks, keywords)
                st.dataframe(df)

    if st.button("📄 PDF 보고서 생성"):
        rules = extract_rules_advanced(text)
        blocks, keywords = extract_sections(text)
        df, keyword_map = analyze_keywords(blocks, keywords)
        os.makedirs("output", exist_ok=True)
        pdf_path = "output/사주_보고서.pdf"
        generate_pdf_report(pdf_path, rules, df)
        st.download_button("📥 다운로드", open(pdf_path, "rb"), file_name="사주_보고서.pdf")
